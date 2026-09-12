import re
import sqlite3
import unicodedata
from io import BytesIO
from urllib.parse import urljoin

import requests
from bs4 import BeautifulSoup
from flask import flash, redirect, render_template, request, url_for
from openpyxl import load_workbook
from pypdf import PdfReader

SOURCES = {
    "aomori": {
        "prefecture": "青森県",
        "title": "住宅宿泊事業法届出住宅一覧",
        "url": "https://opendata.pref.aomori.lg.jp/dataset/2047.html",
        "format": "XLSX（オープンデータ）",
        "mode": "import",
        "note": "県公式オープンデータのXLSXをヘッダー名で判定し、届出住宅所在地を候補へ取り込み。",
    },
    "iwate": {
        "prefecture": "岩手県",
        "title": "民泊（住宅宿泊事業）について",
        "url": "https://www.pref.iwate.jp/kurashikankyou/anzenanshin/seikatsueisei/1016554/1016555.html",
        "format": "PDF（届出状況一覧）",
        "mode": "source_only",
        "note": "公式PDFは届出状況の集計中心。個別住所を誤抽出しないため自動登録は停止。",
    },
    "akita": {
        "prefecture": "秋田県",
        "title": "住宅宿泊事業者一覧",
        "url": "https://www.pref.akita.lg.jp/pages/archive/31592",
        "format": "PDF（届出者・住宅所在地・電話）",
        "mode": "import",
        "note": "県公式PDFの列位置を確認し、届出者住所と届出住宅所在地を区別して候補へ取り込み。",
    },
    "yamagata": {
        "prefecture": "山形県",
        "title": "住宅宿泊事業者一覧",
        "url": "https://www.pref.yamagata.jp/020071/kenfuku/doubutsuaigo/eisei/seikatsueisei/minpaku/minpaku.html",
        "format": "PDF（届出住宅所在地）",
        "mode": "import",
        "note": "県公式PDFが届出住宅所在地を一覧公開しているため、安全に住所候補として取り込み。",
    },
    "fukushima": {
        "prefecture": "福島県",
        "title": "住宅宿泊事業法の受理済み届出住宅一覧",
        "url": "https://www.pref.fukushima.lg.jp/sec/32031a/minpaku-04.html",
        "format": "HTML（商号・屋号・住所等）",
        "mode": "import",
        "note": "県公式HTMLの表から商号・屋号・届出住宅住所・公開電話を取り込み。",
    },
}

CORPORATE_WORDS = ("株式会社", "有限会社", "合同会社", "一般社団法人", "一般財団法人", "（株）", "(株)", "（有）", "(有)")
MAX_BYTES = 10 * 1024 * 1024


def _clean(value):
    return unicodedata.normalize("NFKC", str(value or "")).strip()


def _normalize(value):
    text = unicodedata.normalize("NFKC", value or "").strip().lower()
    text = re.sub(r"[\s\u3000]+", "", text)
    text = re.sub(r"[‐‑‒–—―ー−﹣－]+", "-", text)
    return text


def _candidate_key(prefecture, city, address):
    return "|".join(_normalize(part) for part in (prefecture, city, address))


def _extract_city(address):
    text = (address or "").strip()
    match = re.match(r"^(.+?郡.+?[町村]|.+?[市区町村])", text)
    return match.group(1) if match else ""


def _get(url, accept="text/html,*/*"):
    response = requests.get(url, timeout=20, headers={
        "User-Agent": "Mozilla/5.0 TohokuLodgingSalesAI/1.0",
        "Accept": accept,
        "Accept-Language": "ja,en;q=0.7",
    }, allow_redirects=True)
    response.raise_for_status()
    if len(response.content) > MAX_BYTES:
        raise ValueError("公式資料のサイズが想定上限を超えています。")
    return response


def _discover_resource(page_html, base_url, suffix, label_keyword=""):
    soup = BeautifulSoup(page_html, "html.parser")
    candidates = []
    for link in soup.find_all("a", href=True):
        href = link.get("href", "").strip()
        label = link.get_text(" ", strip=True)
        if suffix.lower() not in href.lower():
            continue
        score = 1
        if label_keyword and label_keyword in label:
            score += 4
        if label_keyword and label_keyword in href:
            score += 2
        candidates.append((score, urljoin(base_url, href)))
    if not candidates:
        raise ValueError(f"公式ページから{suffix.upper()}資料を見つけられませんでした。")
    candidates.sort(key=lambda item: item[0], reverse=True)
    return candidates[0][1]


def discover_aomori_xlsx(page_html, base_url):
    return _discover_resource(page_html, base_url, ".xlsx", "住宅宿泊")


def _header_index(values, aliases):
    normalized = [_normalize(_clean(value)).replace("・", "").replace("、", "") for value in values]
    for index, value in enumerate(normalized):
        if any(_normalize(alias).replace("・", "").replace("、", "") in value for alias in aliases):
            return index
    return None


def parse_aomori_xlsx(xlsx_bytes):
    workbook = load_workbook(BytesIO(xlsx_bytes), read_only=True, data_only=True)
    records = []
    seen = set()
    for sheet in workbook.worksheets:
        rows = list(sheet.iter_rows(values_only=True))
        header_row = None
        mapping = None
        for row_index, row in enumerate(rows[:20]):
            address_index = _header_index(row, ("届出住宅の所在地", "届出住宅所在地", "住宅の所在地", "所在地"))
            if address_index is None:
                continue
            mapping = {
                "address": address_index,
                "name": _header_index(row, ("商号", "屋号", "名称")),
                "operator": _header_index(row, ("届出者氏名又は名称", "届出者氏名", "届出者名称", "氏名又は名称")),
                "phone": _header_index(row, ("電話番号", "連絡先")),
            }
            header_row = row_index
            break
        if mapping is None:
            continue
        for row in rows[header_row + 1:]:
            address = _clean(row[mapping["address"]]) if mapping["address"] < len(row) else ""
            if not address or not re.search(r"[市町村郡]", address):
                continue
            key = _normalize(address)
            if key in seen:
                continue
            seen.add(key)
            name = _clean(row[mapping["name"]]) if mapping["name"] is not None and mapping["name"] < len(row) else ""
            operator = _clean(row[mapping["operator"]]) if mapping["operator"] is not None and mapping["operator"] < len(row) else ""
            phone = _clean(row[mapping["phone"]]) if mapping["phone"] is not None and mapping["phone"] < len(row) else ""
            records.append({
                "address": address,
                "name": name,
                "company_name": operator if any(word in operator for word in CORPORATE_WORDS) else "",
                "phone": phone,
                "official_url": "",
            })
    if not records:
        raise ValueError("青森県XLSXから届出住宅所在地を抽出できませんでした。")
    return records


def fetch_aomori_records():
    page = _get(SOURCES["aomori"]["url"])
    xlsx_url = discover_aomori_xlsx(page.text, SOURCES["aomori"]["url"])
    data = _get(xlsx_url, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet,*/*").content
    if not data.startswith(b"PK"):
        raise ValueError("青森県の取得資料がXLSX形式ではありません。")
    return parse_aomori_xlsx(data), xlsx_url


def discover_akita_pdf(page_html, base_url):
    soup = BeautifulSoup(page_html, "html.parser")
    candidates = []
    for link in soup.find_all("a", href=True):
        href = link.get("href", "").strip()
        label = link.get_text(" ", strip=True)
        if ".pdf" not in href.lower():
            continue
        if "住宅宿泊事業者一覧" not in label and "住宅宿泊事業者一覧" not in href:
            continue
        candidates.append(urljoin(base_url, href))
    if not candidates:
        raise ValueError("秋田県の住宅宿泊事業者一覧PDFを見つけられませんでした。")
    return candidates[-1]


def _layout_text(page):
    try:
        return page.extract_text(extraction_mode="layout") or ""
    except TypeError:
        return page.extract_text() or ""


def parse_akita_pdf(pdf_bytes):
    text = "\n".join(_layout_text(page) for page in PdfReader(BytesIO(pdf_bytes)).pages)
    lines = [unicodedata.normalize("NFKC", line.rstrip()) for line in text.splitlines() if line.strip()]
    header_line = next((line for line in lines if "届出番号" in line and "届出者住所" in line and "届出住宅の所在地" in line), "")
    if not header_line:
        raise ValueError("秋田県PDFの表ヘッダーを確認できませんでした。")

    number_pos = header_line.find("届出番号")
    owner_pos = header_line.find("届出者氏名")
    owner_address_pos = header_line.find("届出者住所")
    residence_pos = header_line.find("届出住宅の所在地")
    phone_pos = header_line.find("電話番号")
    if min(number_pos, owner_pos, owner_address_pos, residence_pos, phone_pos) < 0:
        raise ValueError("秋田県PDFの列位置を確認できませんでした。")

    starts = []
    for index, line in enumerate(lines):
        if re.match(r"^\s*\d+\s+M\d+", line):
            starts.append(index)
    records = []
    seen = set()
    for pos, start in enumerate(starts):
        end = starts[pos + 1] if pos + 1 < len(starts) else len(lines)
        block = lines[start:end]
        owner_parts = []
        owner_address_parts = []
        residence_parts = []
        phone_parts = []
        for line in block:
            padded = line + " " * max(0, phone_pos + 30 - len(line))
            owner_parts.append(padded[owner_pos:owner_address_pos].strip())
            owner_address_parts.append(padded[owner_address_pos:residence_pos].strip())
            residence_parts.append(padded[residence_pos:phone_pos].strip())
            phone_parts.append(padded[phone_pos:].strip())
        owner = " ".join(part for part in owner_parts if part)
        owner_address = " ".join(part for part in owner_address_parts if part)
        residence = " ".join(part for part in residence_parts if part)
        phone_text = " ".join(part for part in phone_parts if part)
        residence = re.sub(r"\s+", " ", residence).strip()
        owner_address = re.sub(r"\s+", " ", owner_address).strip()
        if residence.startswith("同左"):
            residence = owner_address + residence[2:]
        phone_match = re.search(r"0\d{1,4}-\d{1,4}-\d{3,4}", phone_text)
        phone = phone_match.group(0) if phone_match else ""
        if not residence or not re.search(r"[市町村郡]", residence):
            continue
        key = _normalize(residence)
        if key in seen:
            continue
        seen.add(key)
        records.append({
            "address": residence,
            "name": "",
            "company_name": owner if any(word in owner for word in CORPORATE_WORDS) else "",
            "phone": phone,
            "official_url": "",
        })
    if not records:
        raise ValueError("秋田県PDFから届出住宅所在地を抽出できませんでした。")
    return records


def fetch_akita_records():
    page = _get(SOURCES["akita"]["url"])
    pdf_url = discover_akita_pdf(page.text, SOURCES["akita"]["url"])
    data = _get(pdf_url, "application/pdf,*/*").content
    if not data.startswith(b"%PDF"):
        raise ValueError("秋田県の取得資料がPDFではありません。")
    return parse_akita_pdf(data), pdf_url


def discover_yamagata_pdf(page_html, base_url):
    soup = BeautifulSoup(page_html, "html.parser")
    for link in soup.find_all("a", href=True):
        label = link.get_text(" ", strip=True)
        href = link.get("href", "")
        if "住宅宿泊事業者一覧" in label and ".pdf" in href.lower():
            return urljoin(base_url, href)
    raise ValueError("山形県の住宅宿泊事業者一覧PDFを見つけられませんでした。")


def parse_yamagata_pdf(pdf_bytes):
    text = "\n".join(page.extract_text() or "" for page in PdfReader(BytesIO(pdf_bytes)).pages)
    addresses = []
    seen = set()
    for raw in text.splitlines():
        line = unicodedata.normalize("NFKC", raw).strip()
        line = re.sub(r"^\s*\d+[.)]?\s*", "", line)
        if not line or "届出住宅" in line or "住宅宿泊" in line or "現在" in line:
            continue
        if not re.match(r"^(.+?郡.+?[町村]|.+?[市町村])", line):
            continue
        key = _normalize(line)
        if key not in seen:
            seen.add(key)
            addresses.append(line)
    return addresses


def fetch_yamagata_records():
    page = _get(SOURCES["yamagata"]["url"])
    pdf_url = discover_yamagata_pdf(page.text, SOURCES["yamagata"]["url"])
    pdf = _get(pdf_url, "application/pdf,*/*").content
    if not pdf.startswith(b"%PDF"):
        raise ValueError("山形県の取得資料がPDFではありません。")
    return [{"address": address, "name": "", "company_name": "", "phone": "", "official_url": ""} for address in parse_yamagata_pdf(pdf)], pdf_url


def parse_fukushima_html(html):
    soup = BeautifulSoup(html, "html.parser")
    records = []
    for row in soup.select("table tr"):
        cells = [cell.get_text(" ", strip=True) for cell in row.find_all("td")]
        if len(cells) < 3:
            continue
        company, name, address = (cells[0].strip(), cells[1].strip(), cells[2].strip())
        phone = cells[3].strip() if len(cells) > 3 else ""
        if not address or not re.search(r"[市町村郡]", address):
            continue
        records.append({
            "address": address,
            "name": name,
            "company_name": company if any(word in company for word in CORPORATE_WORDS) else "",
            "phone": phone,
            "official_url": "",
        })
    return records


def fetch_fukushima_records():
    response = _get(SOURCES["fukushima"]["url"])
    records = parse_fukushima_html(response.text)
    if not records:
        raise ValueError("福島県の公式一覧から候補を抽出できませんでした。")
    return records, SOURCES["fukushima"]["url"]


def import_records(database, source_key, records, source_url):
    source = SOURCES[source_key]
    prefecture = source["prefecture"]
    new_count = duplicate_count = 0
    with sqlite3.connect(database) as con:
        for record in records:
            address = record["address"].strip()
            city = _extract_city(address)
            key = _candidate_key(prefecture, city, address)
            if con.execute("SELECT id FROM lead_candidates WHERE normalized_key=?", (key,)).fetchone():
                duplicate_count += 1
                continue
            con.execute(
                """INSERT INTO lead_candidates
                   (name,company_name,prefecture,city,address,official_url,phone,source_url,source_type,normalized_key,status,research_status,research_notes)
                   VALUES (?,?,?,?,?,?,?,?,?,?,'pending','unresearched',?)""",
                (
                    record.get("name", ""), record.get("company_name", ""), prefecture, city, address,
                    record.get("official_url", ""), record.get("phone", ""), source_url,
                    f"{prefecture} {source['title']}", key, f"[{prefecture}公式資料から収集] {source_url}",
                ),
            )
            new_count += 1
    return {"total": len(records), "new": new_count, "duplicate": duplicate_count}


def register_regional_sources(app):
    if "regional_sources" in app.view_functions:
        return app

    @app.get("/regional-sources", endpoint="regional_sources")
    def regional_sources():
        return render_template("regional_sources.html", sources=SOURCES)

    @app.post("/regional-sources/import", endpoint="import_regional_source")
    def import_regional_source():
        source_key = request.form.get("source", "")
        source = SOURCES.get(source_key)
        if not source or source["mode"] != "import":
            flash("この公式資料は安全な自動取込の対象外です。公式ページから個別確認してください。", "error")
            return redirect(url_for("regional_sources"))
        try:
            if source_key == "aomori":
                records, source_url = fetch_aomori_records()
            elif source_key == "akita":
                records, source_url = fetch_akita_records()
            elif source_key == "yamagata":
                records, source_url = fetch_yamagata_records()
            elif source_key == "fukushima":
                records, source_url = fetch_fukushima_records()
            else:
                raise ValueError("未対応の公式資料です。")
            stats = import_records(app.config["DATABASE"], source_key, records, source_url)
            flash(
                f"{source['prefecture']}：公式資料{stats['total']}件を確認し、新規{stats['new']}件・重複{stats['duplicate']}件でした。",
                "success",
            )
        except (requests.RequestException, ValueError, OSError) as exc:
            app.logger.warning("Regional source import failed (%s): %s", source_key, exc)
            flash(f"{source['prefecture']}の公式資料を取得・解析できませんでした。時間を置いて再度お試しください。", "error")
        return redirect(url_for("regional_sources"))

    return app
