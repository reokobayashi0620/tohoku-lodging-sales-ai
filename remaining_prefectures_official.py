import io
import re

import requests
from bs4 import BeautifulSoup
from flask import flash, redirect
from openpyxl import load_workbook

from app import download_pdf, pdf_text
from regional_official_import import import_official_records, _clean

IWATE_PAGE = "https://www.pref.iwate.jp/kurashikankyou/anzenanshin/seikatsueisei/1016554/1016555.html"
AKITA_PAGE = "https://www.pref.akita.lg.jp/pages/archive/31592"
AOMORI_PAGE = "https://opendata.pref.aomori.lg.jp/dataset/2047.html"
IWATE_SOURCE_TYPE = "岩手県公式 住宅宿泊事業法届出状況一覧"
AKITA_SOURCE_TYPE = "秋田県公式 住宅宿泊事業者一覧"
AOMORI_SOURCE_TYPE = "青森県公式 住宅宿泊事業法届出住宅一覧"
HEADERS = {"User-Agent": "Mozilla/5.0 TohokuLodgingSalesAI/1.0", "Accept-Language": "ja,en;q=0.7"}


def _latest_link(page_url, keywords, suffix):
    response = requests.get(page_url, timeout=20, headers=HEADERS)
    response.raise_for_status()
    soup = BeautifulSoup(response.text, "html.parser")
    matches = []
    for a in soup.find_all("a", href=True):
        label = _clean(a.get_text(" ", strip=True))
        href = a["href"]
        if all(word in label for word in keywords) and suffix.lower() in href.lower():
            matches.append(requests.compat.urljoin(page_url, href))
    if not matches:
        raise ValueError("公式一覧のダウンロードリンクを確認できませんでした。")
    return matches[-1]


def _pdf_records(url, prefecture, with_operator=False):
    text = pdf_text(download_pdf(url))
    records, seen = [], set()
    for raw in text.splitlines():
        line = _clean(raw)
        if not line or "届出" in line and "所在地" in line:
            continue
        address_match = re.search(r"((?:%s)?[^\s　]{0,20}[市郡][^\s　]{0,30}(?:市|町|村|区)?[^\s　]{0,40})" % prefecture, line)
        if not address_match:
            continue
        address = address_match.group(1).strip()
        if len(address) < 5 or address in seen:
            continue
        seen.add(address)
        operator = ""
        phone = ""
        if with_operator:
            corp = re.search(r"((?:株式会社|有限会社|合同会社|一般社団法人)[^\s　]+)", line)
            operator = corp.group(1) if corp else ""
            tel = re.search(r"0\d{1,4}[-－]\d{1,4}[-－]\d{3,4}", line)
            phone = tel.group(0).replace("－", "-") if tel else ""
        records.append({"name": "", "operator": operator, "address": address, "phone": phone})
    if not records:
        raise ValueError("公式PDFから届出住宅を抽出できませんでした。")
    return records


def fetch_iwate_records():
    url = _latest_link(IWATE_PAGE, ("住宅宿泊事業法", "届出状況一覧"), ".pdf")
    return _pdf_records(url, "岩手県"), url


def fetch_akita_records():
    url = _latest_link(AKITA_PAGE, ("住宅宿泊事業者一覧",), ".pdf")
    return _pdf_records(url, "秋田県", with_operator=True), url


def parse_aomori_xlsx(data):
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    records, seen = [], set()
    for ws in wb.worksheets:
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            continue
        header_index = None
        address_col = None
        for i, row in enumerate(rows[:20]):
            labels = [_clean(str(v)) if v is not None else "" for v in row]
            for j, label in enumerate(labels):
                if "届出住宅" in label and "所在地" in label:
                    header_index, address_col = i, j
                    break
            if address_col is not None:
                break
        if address_col is None:
            continue
        for row in rows[header_index + 1:]:
            if address_col >= len(row) or row[address_col] is None:
                continue
            address = _clean(str(row[address_col]))
            if not address or address in seen or not re.search(r"[市町村]", address):
                continue
            seen.add(address)
            records.append({"name": "", "operator": "", "address": address, "phone": ""})
    if not records:
        raise ValueError("青森県公式XLSXから届出住宅を抽出できませんでした。")
    return records


def fetch_aomori_records():
    url = _latest_link(AOMORI_PAGE, ("住宅宿泊事業届出一覧",), ".xlsx")
    response = requests.get(url, timeout=20, headers=HEADERS)
    response.raise_for_status()
    return parse_aomori_xlsx(response.content), url


def register_remaining_prefectures_official(app):
    configs = {
        "iwate": ("岩手県", fetch_iwate_records, IWATE_SOURCE_TYPE),
        "akita": ("秋田県", fetch_akita_records, AKITA_SOURCE_TYPE),
        "aomori": ("青森県", fetch_aomori_records, AOMORI_SOURCE_TYPE),
    }
    for slug, (prefecture, fetcher, source_type) in configs.items():
        def handler(prefecture=prefecture, fetcher=fetcher, source_type=source_type):
            try:
                records, source_url = fetcher()
                stats = import_official_records(app.config["DATABASE"], prefecture, records, source_url, source_type)
                flash(f"{prefecture}公式民泊一覧 {stats['source_records']}件を確認。新規{stats['inserted']}件、既存更新{stats['updated']}件です。", "success")
            except (requests.RequestException, ValueError, OSError) as exc:
                app.logger.warning("%s official import failed: %s", prefecture, exc)
                flash(f"{prefecture}公式民泊一覧を取得できませんでした。時間を置いて再度お試しください。", "error")
            return redirect(f"/targets?prefecture={prefecture}")
        app.add_url_rule(f"/targets/import-{slug}", endpoint=f"import_{slug}_targets", view_func=handler, methods=["POST"])
    return app
